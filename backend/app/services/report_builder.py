"""
Builds polished PowerPoint, Word and PDF reports from the same analytics data.
Charts are rendered with matplotlib so all three formats look identical.
"""
from __future__ import annotations

import io
from datetime import date

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

from . import analytics  # noqa: E402
from .ai_analyst import write_narrative  # noqa: E402

INK = "#14213D"
AMBER = "#E9A23B"
TEAL = "#2A9D8F"
MUTED = "#6B7A90"
LIGHT = "#EEF2F6"
REGION_COLORS = {"Europe": "#2A9D8F", "North America": "#14213D", "Asia Pacific": "#E9A23B",
                 "Middle East": "#C8553D", "Latin America": "#8E7DBE", "Africa": "#6B7A90"}

ALL_SECTIONS = ["summary", "kpis", "trend", "markets", "retention", "funnel", "expansion", "recommendations"]


# ---------------------------------------------------------------- data + narrative

def build_context(months: int, region: str | None, sections: list[str], weights: dict | None) -> dict:
    ov = analytics.overview(months, region)
    factors = analytics.market_factors()
    w = weights or {"market_size": 35, "digital_adoption": 25, "search_demand": 25, "affluence": 15}
    total_w = sum(w.values()) or 1
    candidates = []
    for f in factors:
        if f["launched"]:
            continue
        score = sum(f[k] * v for k, v in w.items()) / total_w
        candidates.append({"name": f["name"], "region": f["region"], "score": round(score, 1),
                           "search_demand": f["search_demand"], "digital_adoption": f["digital_adoption"]})
    candidates.sort(key=lambda c: c["score"], reverse=True)
    launched = [c for c in ov["countries"] if c["launched"]]
    ctx = {
        "months": months, "region": region or "All regions", "sections": sections,
        "kpis": ov["kpis"], "trend": ov["trend"], "countries": launched, "funnel": ov["funnel"],
        "channels": ov["channels"], "cohorts": analytics.cohorts(region), "expansion": candidates,
    }
    ctx["narrative"] = write_narrative({k: ctx[k] for k in ("kpis", "countries", "funnel", "channels", "expansion")}) \
        or _template_narrative(ctx)
    return ctx


def _template_narrative(ctx: dict) -> dict:
    k = ctx["kpis"]
    c = ctx["countries"]
    fastest = max(c, key=lambda x: x["growth"] or -999) if c else None
    weakest_conv = min(c, key=lambda x: x["conversion"] or 99) if c else None
    top_share = sum(x["share"] for x in c[:3]) if c else 0
    exp = ctx["expansion"][0] if ctx["expansion"] else None
    summary = (
        f"Revenue reached ${k['revenue']['value']:,.0f} over the last {ctx['months']} months, "
        f"{_pct(k['revenue']['change'])} versus the prior period, on {k['orders']['value']:,} orders. "
        f"The top three markets account for {top_share:.0f}% of revenue"
        + (f", and {fastest['name']} is growing fastest at {_pct(fastest['growth'])}." if fastest else ".")
        + (f" {exp['name']} is the most attractive market not yet served." if exp else "")
    )
    insights = []
    if c:
        insights.append(f"{c[0]['name']} is the largest market with ${c[0]['revenue']:,.0f} ({c[0]['share']}% share).")
    if fastest:
        insights.append(f"{fastest['name']} grew {_pct(fastest['growth'])}, the strongest growth in the portfolio.")
    if weakest_conv:
        insights.append(f"{weakest_conv['name']} converts at {weakest_conv['conversion']}% against "
                        f"{k['conversion']['value']}% overall, with friction between cart and checkout.")
    insights.append(f"Average order value is ${k['aov']['value']:.2f} and repeat orders make up "
                    f"{k['repeat_rate']['value']}% of all orders.")
    recs = [
        f"Prioritise growth budget for {fastest['name'] if fastest else 'the fastest-growing markets'}; watch CAC payback monthly.",
        f"Fix checkout friction in {weakest_conv['name'] if weakest_conv else 'low-converting markets'} with local payment "
        "methods and upfront delivery costs; watch cart-to-checkout rate.",
        f"Run a 90-day pilot in {exp['name'] if exp else 'the top expansion candidate'} through a marketplace before "
        "local fulfilment; watch first-order CAC and 60-day repeat rate.",
    ]
    return {"summary": summary, "insights": insights, "recommendations": recs}


def _pct(v) -> str:
    return "n/a" if v is None else f"{v:+.1f}%"


def _money(v) -> str:
    return f"${v / 1e6:.2f}M" if v >= 1e6 else f"${v / 1e3:.0f}K" if v >= 1e3 else f"${v:,.0f}"


def kpi_tiles(ctx: dict) -> list[tuple[str, str, str]]:
    k = ctx["kpis"]
    return [
        ("Revenue", _money(k["revenue"]["value"]), _pct(k["revenue"]["change"])),
        ("Orders", f"{k['orders']['value']:,}", _pct(k["orders"]["change"])),
        ("Avg order value", f"${k['aov']['value']:.2f}", _pct(k["aov"]["change"])),
        ("Repeat orders", f"{k['repeat_rate']['value']}%", _pct(k["repeat_rate"]["change"])),
        ("Conversion", f"{k['conversion']['value']}%", _pct(k["conversion"]["change"])),
        ("Customers", f"{k['customers']['value']:,}", _pct(k["customers"]["change"])),
    ]


# ---------------------------------------------------------------- charts

def _style(ax, title: str):
    ax.set_title(title, loc="left", fontsize=12, color=INK, fontweight="bold", pad=12)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    for s in ("left", "bottom"):
        ax.spines[s].set_color("#C9D2DE")
    ax.tick_params(colors=MUTED, labelsize=8)
    ax.grid(axis="y", color="#E3E8EF", linewidth=0.8)
    ax.set_axisbelow(True)


def _png(fig) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=200, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return buf.getvalue()


def chart_trend(ctx) -> bytes:
    trend = ctx["trend"]
    months = [t["month"] for t in trend]
    fig, ax = plt.subplots(figsize=(8, 3.6))
    bottom = [0] * len(trend)
    for region, color in REGION_COLORS.items():
        vals = [t.get(region, 0) / 1000 for t in trend]
        if sum(vals) == 0:
            continue
        ax.bar(months, vals, bottom=bottom, color=color, label=region, width=0.75)
        bottom = [b + v for b, v in zip(bottom, vals)]
    _style(ax, "Monthly revenue by region (USD thousands)")
    ax.set_xticks(months[::3])
    ax.legend(frameon=False, fontsize=7, ncol=6, loc="upper left", bbox_to_anchor=(0, -0.12))
    return _png(fig)


def chart_markets(ctx) -> bytes:
    top = ctx["countries"][:10][::-1]
    fig, ax = plt.subplots(figsize=(8, 3.8))
    ax.barh([c["name"] for c in top], [c["revenue"] / 1000 for c in top], color=TEAL, height=0.6)
    for i, c in enumerate(top):
        ax.text(c["revenue"] / 1000, i, f"  {_pct(c['growth'])}", va="center", fontsize=7,
                color=AMBER if (c["growth"] or 0) >= 0 else "#C8553D")
    _style(ax, "Top markets by revenue (USD thousands) with growth")
    ax.grid(axis="x", color="#E3E8EF")
    ax.grid(axis="y", visible=False)
    return _png(fig)


def chart_retention(ctx) -> bytes:
    rows = ctx["cohorts"]
    data = [r["retention"][1:] for r in rows]
    fig, ax = plt.subplots(figsize=(8, 3.8))
    im = ax.imshow(data, cmap="YlGnBu", aspect="auto", vmin=0, vmax=max(max(d) for d in data) or 1)
    ax.set_yticks(range(len(rows)), [r["cohort"] for r in rows])
    ax.set_xticks(range(len(data[0])), [f"M{i + 1}" for i in range(len(data[0]))])
    for i, row in enumerate(data):
        for j, v in enumerate(row):
            ax.text(j, i, f"{v:.0f}%", ha="center", va="center", fontsize=6.5,
                    color="white" if v > 15 else INK)
    ax.set_title("Customer retention by first-order cohort", loc="left", fontsize=12, color=INK,
                 fontweight="bold", pad=12)
    ax.tick_params(colors=MUTED, labelsize=7)
    fig.colorbar(im, ax=ax, fraction=0.03)
    return _png(fig)


def chart_funnel(ctx) -> bytes:
    f = ctx["funnel"]
    fig, ax = plt.subplots(figsize=(8, 3))
    labels = [s["step"] for s in f][::-1]
    vals = [s["pct"] for s in f][::-1]
    ax.barh(labels, vals, color=[INK, TEAL, AMBER, "#C9D2DE"][::-1][: len(vals)], height=0.6)
    for i, v in enumerate(vals):
        ax.text(v, i, f"  {v:.2f}%", va="center", fontsize=8, color=INK)
    ax.set_xscale("log")
    _style(ax, "Conversion funnel (% of sessions, log scale)")
    ax.grid(axis="y", visible=False)
    return _png(fig)


def chart_expansion(ctx) -> bytes:
    exp = ctx["expansion"][::-1]
    fig, ax = plt.subplots(figsize=(8, 2.8))
    ax.barh([e["name"] for e in exp], [e["score"] for e in exp], color=AMBER, height=0.55)
    for i, e in enumerate(exp):
        ax.text(e["score"], i, f"  {e['score']}", va="center", fontsize=8, color=INK)
    ax.set_xlim(0, 100)
    _style(ax, "Expansion attractiveness score (0-100)")
    ax.grid(axis="y", visible=False)
    return _png(fig)


CHARTS = {"trend": chart_trend, "markets": chart_markets, "retention": chart_retention,
          "funnel": chart_funnel, "expansion": chart_expansion}

SECTION_TITLES = {
    "summary": "Executive summary", "kpis": "Key metrics", "trend": "Revenue trend",
    "markets": "Market performance", "retention": "Customer retention", "funnel": "Conversion funnel",
    "expansion": "Expansion opportunities", "recommendations": "Recommendations",
}

SECTION_NOTES = {
    "trend": "Revenue by region per month, showing growth and seasonality across the portfolio.",
    "markets": "Largest markets by revenue in the period, labelled with growth against the prior period.",
    "retention": "Share of each monthly cohort of new customers that ordered again in later months.",
    "funnel": "How sessions turn into purchases. The biggest drop shows where to focus optimisation.",
    "expansion": "Markets not yet served, scored on size, digital adoption, search demand and affluence.",
}


# ---------------------------------------------------------------- PowerPoint

def build_pptx(ctx: dict, title: str) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    def rgb(h):
        return RGBColor.from_string(h.lstrip("#"))

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def text(slide, x, y, w, h, s, size=14, bold=False, color=INK, align=None):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = s
        p.font.size, p.font.bold, p.font.color.rgb, p.font.name = Pt(size), bold, rgb(color), "Calibri"
        if align:
            p.alignment = align
        return tf

    def header(slide, heading, sub=None):
        bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.12))
        bar.fill.solid(); bar.fill.fore_color.rgb = rgb(AMBER); bar.line.fill.background()
        text(slide, 0.6, 0.4, 12, 0.8, heading, 28, True)
        if sub:
            text(slide, 0.6, 1.1, 12, 0.5, sub, 13, color=MUTED)
        text(slide, 0.6, 7.0, 6, 0.3, f"MarketLens  |  {ctx['region']}  |  last {ctx['months']} months", 9, color=MUTED)

    # Title slide
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = rgb(INK); bg.line.fill.background()
    accent = s.shapes.add_shape(1, Inches(0.8), Inches(3.55), Inches(1.2), Inches(0.08))
    accent.fill.solid(); accent.fill.fore_color.rgb = rgb(AMBER); accent.line.fill.background()
    text(s, 0.8, 2.2, 11, 1.3, title, 40, True, "#FFFFFF")
    text(s, 0.8, 3.8, 11, 0.6, f"{ctx['region']}  |  last {ctx['months']} months  |  {date.today():%d %B %Y}", 16, color="#C9D2DE")
    text(s, 0.8, 6.6, 11, 0.4, "Prepared with MarketLens", 11, color="#8EA0B8")

    for sec in ctx["sections"]:
        if sec == "summary":
            s = prs.slides.add_slide(blank)
            header(s, "Executive summary")
            text(s, 0.6, 1.5, 12, 1.6, ctx["narrative"]["summary"], 18)
            tf = text(s, 0.6, 3.6, 12, 3, "Key insights", 16, True, TEAL)
            for ins in ctx["narrative"]["insights"]:
                p = tf.add_paragraph(); p.text = f"•  {ins}"; p.font.size = Pt(14); p.font.color.rgb = rgb(INK)
                p.space_before = Pt(8)
        elif sec == "kpis":
            s = prs.slides.add_slide(blank)
            header(s, "Key metrics", f"Change versus the previous {ctx['months']} months")
            for i, (label, val, chg) in enumerate(kpi_tiles(ctx)):
                x, y = 0.6 + (i % 3) * 4.1, 1.8 + (i // 3) * 2.5
                card = s.shapes.add_shape(5, Inches(x), Inches(y), Inches(3.8), Inches(2.1))
                card.fill.solid(); card.fill.fore_color.rgb = rgb(LIGHT); card.line.fill.background()
                text(s, x + 0.3, y + 0.25, 3.3, 0.4, label, 13, color=MUTED)
                text(s, x + 0.3, y + 0.7, 3.3, 0.8, val, 32, True)
                text(s, x + 0.3, y + 1.5, 3.3, 0.4, chg, 13, True, TEAL if not chg.startswith("-") else "#C8553D")
        elif sec in CHARTS:
            s = prs.slides.add_slide(blank)
            header(s, SECTION_TITLES[sec], SECTION_NOTES[sec])
            png = CHARTS[sec](ctx)
            from PIL import Image as PILImage
            iw, ih = PILImage.open(io.BytesIO(png)).size
            w = min(11.5, 5.1 * iw / ih)  # fit inside an 11.5 x 5.1 inch box
            h = w * ih / iw
            s.shapes.add_picture(io.BytesIO(png), Inches((13.333 - w) / 2), Inches(1.7 + (5.1 - h) / 2), width=Inches(w))
        elif sec == "recommendations":
            s = prs.slides.add_slide(blank)
            header(s, "Recommendations")
            for i, rec in enumerate(ctx["narrative"]["recommendations"]):
                y = 1.6 + i * 1.75
                num = s.shapes.add_shape(9, Inches(0.7), Inches(y), Inches(0.7), Inches(0.7))
                num.fill.solid(); num.fill.fore_color.rgb = rgb(AMBER); num.line.fill.background()
                num.text_frame.text = str(i + 1)
                num.text_frame.paragraphs[0].font.size = Pt(18)
                num.text_frame.paragraphs[0].font.bold = True
                num.text_frame.paragraphs[0].font.color.rgb = rgb(INK)
                text(s, 1.7, y + 0.05, 10.8, 1.4, rec, 17)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- Word

def build_docx(ctx: dict, title: str) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    def rgb(h):
        return RGBColor.from_string(h.lstrip("#"))

    doc = Document()
    sec0 = doc.sections[0]
    sec0.left_margin = sec0.right_margin = Inches(0.9)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    for lvl, size in ((1, 18), (2, 13)):
        hs = doc.styles[f"Heading {lvl}"]
        hs.font.name, hs.font.size, hs.font.color.rgb = "Calibri", Pt(size), rgb(INK)

    def shade(cell, hex_color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_color.lstrip("#"))
        tcPr.append(shd)

    t = doc.add_paragraph()
    r = t.add_run(title); r.bold = True; r.font.size = Pt(26); r.font.color.rgb = rgb(INK)
    sub = doc.add_paragraph()
    r = sub.add_run(f"{ctx['region']}  |  last {ctx['months']} months  |  {date.today():%d %B %Y}")
    r.font.color.rgb = rgb(MUTED)

    for sec in ctx["sections"]:
        if sec == "summary":
            doc.add_heading("Executive summary", 1)
            doc.add_paragraph(ctx["narrative"]["summary"])
            doc.add_heading("Key insights", 2)
            for ins in ctx["narrative"]["insights"]:
                doc.add_paragraph(ins, style="List Bullet")
        elif sec == "kpis":
            doc.add_heading("Key metrics", 1)
            tiles = kpi_tiles(ctx)
            table = doc.add_table(rows=1, cols=3)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr = table.rows[0].cells
            for i, h in enumerate(("Metric", "Value", "Change vs prior period")):
                hdr[i].text = h; shade(hdr[i], INK)
                hdr[i].paragraphs[0].runs[0].font.color.rgb = rgb("#FFFFFF")
                hdr[i].paragraphs[0].runs[0].bold = True
            for i, (label, val, chg) in enumerate(tiles):
                cells = table.add_row().cells
                cells[0].text, cells[1].text, cells[2].text = label, val, chg
                if i % 2:
                    for c in cells:
                        shade(c, LIGHT)
        elif sec in CHARTS:
            doc.add_heading(SECTION_TITLES[sec], 1)
            doc.add_paragraph(SECTION_NOTES[sec]).runs[0].font.color.rgb = rgb(MUTED)
            doc.add_picture(io.BytesIO(CHARTS[sec](ctx)), width=Inches(6.6))
            if sec == "markets":
                table = doc.add_table(rows=1, cols=5)
                for i, h in enumerate(("Market", "Revenue", "Growth", "Repeat orders", "Conversion")):
                    c = table.rows[0].cells[i]; c.text = h; shade(c, INK)
                    c.paragraphs[0].runs[0].font.color.rgb = rgb("#FFFFFF"); c.paragraphs[0].runs[0].bold = True
                for i, m in enumerate(ctx["countries"][:12]):
                    cells = table.add_row().cells
                    vals = (m["name"], _money(m["revenue"]), _pct(m["growth"]), f"{m['repeat_rate']}%", f"{m['conversion']}%")
                    for c, v in zip(cells, vals):
                        c.text = v
                        if i % 2:
                            shade(c, LIGHT)
        elif sec == "recommendations":
            doc.add_heading("Recommendations", 1)
            for rec in ctx["narrative"]["recommendations"]:
                doc.add_paragraph(rec, style="List Number")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PDF

def build_pdf(ctx: dict, title: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (Image, KeepTogether, ListFlowable, ListItem, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm, title=title)
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], textColor=colors.HexColor(INK), fontSize=17, spaceBefore=14, spaceAfter=8)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], textColor=colors.HexColor(TEAL), fontSize=12, spaceBefore=8)
    body = ParagraphStyle("b", parent=ss["BodyText"], fontSize=10.5, leading=15, textColor=colors.HexColor("#1F2A3D"))
    note = ParagraphStyle("n", parent=body, textColor=colors.HexColor(MUTED), fontSize=9.5)
    big = ParagraphStyle("t", parent=ss["Title"], alignment=0, textColor=colors.HexColor(INK), fontSize=26, leading=31)

    def footer(canvas, d):
        canvas.saveState()
        canvas.setFillColor(colors.HexColor(AMBER))
        canvas.rect(0, A4[1] - 6, A4[0], 6, stroke=0, fill=1)
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor(MUTED))
        canvas.drawString(2 * cm, 1.2 * cm, f"MarketLens  |  {ctx['region']}  |  last {ctx['months']} months")
        canvas.drawRightString(A4[0] - 2 * cm, 1.2 * cm, f"Page {d.page}")
        canvas.restoreState()

    def table_style(n):
        st = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(INK)),
              ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
              ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
              ("FONTSIZE", (0, 0), (-1, -1), 9),
              ("BOTTOMPADDING", (0, 0), (-1, -1), 6), ("TOPPADDING", (0, 0), (-1, -1), 6)]
        st += [("BACKGROUND", (0, i), (-1, i), colors.HexColor(LIGHT)) for i in range(2, n, 2)]
        return TableStyle(st)

    flow = [Paragraph(title, big),
            Paragraph(f"{ctx['region']}  |  last {ctx['months']} months  |  {date.today():%d %B %Y}", note),
            Spacer(1, 10)]
    width = A4[0] - 4 * cm
    for sec in ctx["sections"]:
        if sec == "summary":
            flow += [Paragraph("Executive summary", h1), Paragraph(ctx["narrative"]["summary"], body),
                     Paragraph("Key insights", h2),
                     ListFlowable([ListItem(Paragraph(i, body)) for i in ctx["narrative"]["insights"]],
                                  bulletType="bullet", bulletColor=colors.HexColor(AMBER))]
        elif sec == "kpis":
            data = [["Metric", "Value", "Change vs prior period"]] + [list(t) for t in kpi_tiles(ctx)]
            tbl = Table(data, colWidths=[width * 0.4, width * 0.3, width * 0.3])
            tbl.setStyle(table_style(len(data)))
            flow.append(KeepTogether([Paragraph("Key metrics", h1), tbl]))
        elif sec in CHARTS:
            png = CHARTS[sec](ctx)
            from reportlab.lib.utils import ImageReader
            iw, ih = ImageReader(io.BytesIO(png)).getSize()
            flow.append(KeepTogether([Paragraph(SECTION_TITLES[sec], h1), Paragraph(SECTION_NOTES[sec], note),
                                      Spacer(1, 6), Image(io.BytesIO(png), width=width, height=width * ih / iw)]))
            if sec == "markets":
                data = [["Market", "Revenue", "Growth", "Repeat", "Conversion"]] + [
                    [m["name"], _money(m["revenue"]), _pct(m["growth"]), f"{m['repeat_rate']}%", f"{m['conversion']}%"]
                    for m in ctx["countries"][:12]]
                tbl = Table(data, colWidths=[width * 0.32] + [width * 0.17] * 4)
                tbl.setStyle(table_style(len(data)))
                flow += [Spacer(1, 8), tbl]
        elif sec == "recommendations":
            flow += [Paragraph("Recommendations", h1),
                     ListFlowable([ListItem(Paragraph(r, body)) for r in ctx["narrative"]["recommendations"]],
                                  bulletType="1", bulletColor=colors.HexColor(INK))]
    doc.build(flow, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


BUILDERS = {
    "pptx": (build_pptx, "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
    "docx": (build_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "pdf": (build_pdf, "application/pdf"),
}
